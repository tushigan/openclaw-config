from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _split_lines(value: str | None) -> list[str]:
    if not value:
        return []
    return [line.strip() for line in value.splitlines() if line.strip()]


def _safe_id_part(value: str | None, fallback: str) -> str:
    return value or fallback


def _build_header_map(worksheet: Any) -> dict[str, int]:
    header_positions: dict[str, list[int]] = {}

    for column_index in range(1, worksheet.max_column + 1):
        header = _clean(worksheet.cell(row=1, column=column_index).value)
        if header is None:
            continue
        header_positions.setdefault(header, []).append(column_index)

    required_headers = [
        "项目",
        "任务",
        "说明",
        "服务部门",
        "计划时间(工作日)",
        "交付成果",
        "建议报价",
    ]
    missing_headers = [header for header in required_headers if header not in header_positions]
    if missing_headers:
        raise ValueError(f"汇总表缺少必要表头: {', '.join(missing_headers)}")

    module_columns = header_positions.get("模块", [])
    if len(module_columns) < 2:
        raise ValueError("汇总表需要两列“模块”表头")

    if len(module_columns) > 2:
        counts = Counter(module_columns)
        duplicates = [str(column) for column, count in counts.items() if count > 1]
        if duplicates:
            raise ValueError(f"模块表头定位异常: {', '.join(duplicates)}")

    return {
        "level_1_module": module_columns[0],
        "level_2_module": module_columns[1],
        "project": header_positions["项目"][0],
        "task": header_positions["任务"][0],
        "description": header_positions["说明"][0],
        "service_department": header_positions["服务部门"][0],
        "planned_workdays": header_positions["计划时间(工作日)"][0],
        "deliverables": header_positions["交付成果"][0],
        "suggested_price": header_positions["建议报价"][0],
    }


def parse_summary_sheet(workbook_path: Path, sheet_name: str = "汇总表") -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook[sheet_name]
    header_map = _build_header_map(worksheet)

    current_parents: dict[str, str | None] = {
        "level_1_module": None,
        "level_2_module": None,
        "project": None,
    }
    services: list[dict[str, Any]] = []

    for row_index in range(2, worksheet.max_row + 1):
        parent_row = {
            "level_1_module": _clean(worksheet.cell(row=row_index, column=header_map["level_1_module"]).value),
            "level_2_module": _clean(worksheet.cell(row=row_index, column=header_map["level_2_module"]).value),
            "project": _clean(worksheet.cell(row=row_index, column=header_map["project"]).value),
        }
        row = {
            "task": _clean(worksheet.cell(row=row_index, column=header_map["task"]).value),
            "description": _clean(worksheet.cell(row=row_index, column=header_map["description"]).value),
            "service_department": _clean(worksheet.cell(row=row_index, column=header_map["service_department"]).value),
            "planned_workdays": _clean(worksheet.cell(row=row_index, column=header_map["planned_workdays"]).value),
            "deliverables": _split_lines(
                _clean(worksheet.cell(row=row_index, column=header_map["deliverables"]).value)
            ),
            "suggested_price": _clean(worksheet.cell(row=row_index, column=header_map["suggested_price"]).value),
        }

        if not any(value not in (None, [], "") for value in row.values()):
            continue

        if row["task"] is None:
            continue

        for key, value in parent_row.items():
            if value is not None:
                current_parents[key] = value

        service = {
            "service_id": "__".join(
                [
                    _safe_id_part(current_parents["level_1_module"], "unknown-l1"),
                    _safe_id_part(current_parents["level_2_module"], "unknown-l2"),
                    _safe_id_part(current_parents["project"], "unknown-project"),
                    _safe_id_part(row["task"], f"row-{row_index}"),
                ]
            ),
            "level_1_module": current_parents["level_1_module"],
            "level_2_module": current_parents["level_2_module"],
            "project": current_parents["project"],
            "task": row["task"],
            "description": row["description"],
            "service_department": row["service_department"],
            "planned_workdays": row["planned_workdays"],
            "deliverables": row["deliverables"],
            "suggested_price": row["suggested_price"],
            "source": {"sheet": sheet_name, "row": row_index},
        }
        services.append(service)

    return services


def build_catalog(workbook_path: Path, output_path: Path, sheet_name: str = "汇总表") -> dict[str, Any]:
    services = parse_summary_sheet(workbook_path, sheet_name=sheet_name)
    payload = {
        "catalog_name": workbook_path.stem,
        "source_file": str(workbook_path),
        "source_sheet": sheet_name,
        "generated_at": datetime.now().astimezone().isoformat(),
        "services": services,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload
