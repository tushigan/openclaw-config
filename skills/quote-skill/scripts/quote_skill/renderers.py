from __future__ import annotations

import os
import shutil
from pathlib import Path
import subprocess
from string import Template
import tempfile
from typing import Any
from html import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from quote_skill.layout import build_stage_table_sections
from quote_skill.validators import validate_quote_content

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_HTML_TEMPLATE_PATH = REPO_ROOT / "assets" / "templates" / "quote.html"
DEFAULT_HTML_TO_PDF_SCRIPT = REPO_ROOT / "scripts" / "render_html_to_pdf.mjs"
DEFAULT_NODE_BIN = Path(shutil.which("node") or shutil.which("nodejs") or "node")

HEADER_FILL = PatternFill(fill_type="solid", fgColor="F2E7DA")
HEADER_FONT = Font(bold=True, color="4B3A30")
SECTION_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=18)
SUBTITLE_FONT = Font(size=12, color="6B6057")
WRAP_TOP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "；".join(_as_text(item) for item in value if item is not None)
    return str(value)


def _first_present(item: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = item.get(key)
        if value not in (None, "", []):
            return _as_text(value)
    return ""


def _paragraphs_html(paragraphs: list[str]) -> str:
    return "".join(f"<p>{escape(_as_text(paragraph))}</p>" for paragraph in paragraphs)


def _pricing_summary_html(pricing_summary: dict[str, Any]) -> str:
    if not pricing_summary:
        return ""
    return (
        '<div class="note">'
        f"<strong>{escape(_as_text(pricing_summary.get('label')))}</strong><br/>"
        f"{escape(_as_text(pricing_summary.get('details')))}<br/>"
        f"{escape(_as_text(pricing_summary.get('total')))}"
        "</div>"
    )


def _stage_plan_html(stage_plan: list[dict[str, Any]]) -> str:
    if not stage_plan:
        return ""
    header = (
        '<section class="section">'
        "<h2>项目阶段规划</h2>"
        "<table>"
        "<thead><tr>"
        "<th>阶段</th><th>对应板块</th><th>费用</th><th>启动条件</th><th>预估周期</th><th>阶段核心交付</th>"
        "</tr></thead><tbody>"
    )
    rows = []
    for item in stage_plan:
        rows.append(
            "<tr>"
            f"<td>{escape(_first_present(item, 'stage', 'name'))}</td>"
            f"<td>{escape(_first_present(item, 'module'))}</td>"
            f"<td>{escape(_first_present(item, 'price'))}</td>"
            f"<td>{escape(_first_present(item, 'start_condition'))}</td>"
            f"<td>{escape(_first_present(item, 'duration', 'timeline'))}</td>"
            f"<td>{escape(_first_present(item, 'deliverables'))}</td>"
            "</tr>"
        )
    return header + "".join(rows) + "</tbody></table></section>"


def _legacy_service_sections_html(service_sections: list[dict[str, Any]]) -> str:
    if not service_sections:
        return ""

    sections: list[str] = []
    for section in service_sections:
        title = escape(_as_text(section.get("title")))
        rows = section.get("rows")
        items = section.get("items")
        if rows:
            body_rows = "".join(
                "<tr>"
                f"<td>{escape(_first_present(row, 'block'))}</td>"
                f"<td>{escape(_first_present(row, 'project'))}</td>"
                f"<td>{escape(_first_present(row, 'task'))}</td>"
                f"<td>{escape(_first_present(row, 'description_bullets'))}</td>"
                f"<td>{escape(_first_present(row, 'deliverable'))}</td>"
                f"<td>{escape(_first_present(row, 'department'))}</td>"
                f"<td>{escape(_first_present(row, 'owner'))}</td>"
                "</tr>"
                for row in rows
            )
            sections.append(
                '<section class="section">'
                f"<h2>{title}</h2>"
                "<table><thead><tr>"
                "<th>版块</th><th>项目</th><th>任务</th><th>说明</th><th>交付成果</th><th>服务部门</th><th>负责人</th>"
                f"</tr></thead><tbody>{body_rows}</tbody></table></section>"
            )
            continue

        list_items = "".join(f"<li>{escape(_as_text(item))}</li>" for item in items or [])
        sections.append(
            '<section class="section">'
            f"<h2>{title}</h2>"
            f"<ul>{list_items}</ul>"
            "</section>"
        )
    return "".join(sections)


def _closing_notes_html(notes: list[str]) -> str:
    if not notes:
        return ""
    body = "".join(f'<div class="note">{escape(_as_text(note))}</div>' for note in notes)
    return f'<section class="section"><h2>备注说明</h2>{body}</section>'


def _pricing_summary_lines(pricing_summary: dict[str, Any]) -> list[str]:
    if not pricing_summary:
        return []

    lines = []
    label = _as_text(pricing_summary.get("label"))
    details = _as_text(pricing_summary.get("details"))
    total = _as_text(pricing_summary.get("total"))

    if label:
        lines.append(label)
    if details:
        lines.append(details)
    if total:
        lines.append(total)
    return lines


def _uses_stages(content: dict[str, Any]) -> bool:
    return bool(content.get("stages"))


def _normalized_content(content: dict[str, Any]) -> dict[str, Any]:
    if _uses_stages(content):
        return validate_quote_content(content)
    return content


def _stage_summary_html(section: dict[str, Any]) -> str:
    details = [
        f"阶段费用：{escape(_as_text(section.get('stage_price')) or '待补充')}",
        f"阶段周期：{escape(_as_text(section.get('stage_duration')) or '待确认')}",
    ]
    deliverables = _as_text(section.get("stage_deliverables_text"))
    if deliverables:
        details.append(f"阶段交付：{escape(deliverables)}")
    return f'<div class="small">{" | ".join(details)}</div>'


def _top_level_label(content: dict[str, Any]) -> str:
    if _as_text(content.get("case_type")) == "全案":
        return "服务模块"
    for stage in content.get("stages", []):
        if _as_text(stage.get("stage_kind")) == "service_module":
            return "服务模块"
    return "阶段"


def _staged_sections_html(content: dict[str, Any]) -> str:
    sections = build_stage_table_sections(content)
    html_sections: list[str] = []
    top_level_label = _top_level_label(content)

    for section in sections:
        body_rows: list[str] = []
        for row in section["rows"]:
            cells: list[str] = []
            if row["show_stage"]:
                cells.append(
                    f'<td rowspan="{row["stage_rowspan"]}">{escape(row["stage_name"])}</td>'
                )
            if row["show_project"]:
                cells.append(
                    f'<td rowspan="{row["project_rowspan"]}">{escape(row["project_name"])}</td>'
                )
            cells.append(f"<td>{escape(row['task_name'])}</td>")
            cells.append(f"<td>{escape(row['description_text'])}</td>")
            if row["show_deliverables"]:
                cells.append(
                    f'<td rowspan="{row["project_rowspan"]}">{escape(row["deliverables_text"])}</td>'
                )
            cells.append(f"<td>{escape(row['department_text'])}</td>")
            cells.append(f"<td>{escape(row['owner_text'])}</td>")
            body_rows.append(f"<tr>{''.join(cells)}</tr>")

        html_sections.append(
            '<section class="section">'
            f"<h2>{escape(section['stage_name'])}</h2>"
            f"{_stage_summary_html(section)}"
            "<table><thead><tr>"
            f"<th>{top_level_label}</th><th>项目</th><th>任务</th><th>说明</th><th>交付物</th><th>服务部门</th><th>负责人</th>"
            f"</tr></thead><tbody>{''.join(body_rows)}</tbody></table></section>"
        )

    return "".join(html_sections)


def render_html(
    content: dict[str, Any],
    template_path: Path,
    output_path: Path | None = None,
) -> str:
    content = _normalized_content(content)
    template = Template(template_path.read_text(encoding="utf-8"))
    html = template.safe_substitute(
        title=_as_text(content.get("title")),
        subtitle=_as_text(content.get("subtitle")),
        overview_html=_paragraphs_html(content.get("overview_paragraphs", [])),
        pricing_summary_html=_pricing_summary_html(content.get("pricing_summary", {})),
        stage_plan_html="" if _uses_stages(content) else _stage_plan_html(content.get("stage_plan", [])),
        service_sections_html=(
            _staged_sections_html(content)
            if _uses_stages(content)
            else _legacy_service_sections_html(content.get("service_sections", []))
        ),
        closing_notes_html=_closing_notes_html(content.get("closing_notes", [])),
    )
    if output_path is not None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(html, encoding="utf-8")
    return html


def _render_xlsx_legacy(
    content: dict[str, Any],
    output_path: Path,
    template_path: Path | None = None,
) -> Path:
    if template_path and template_path.exists():
        workbook = load_workbook(template_path)
        sheet = workbook[workbook.sheetnames[0]]
    else:
        workbook = Workbook()
        sheet = workbook.active

    sheet.title = "报价方案"
    sheet["A1"] = _as_text(content.get("title"))
    sheet["A2"] = _as_text(content.get("subtitle"))

    row_cursor = 4
    sheet[f"A{row_cursor}"] = "项目概述"
    for paragraph in content.get("overview_paragraphs", []):
        row_cursor += 1
        sheet[f"A{row_cursor}"] = _as_text(paragraph)

    pricing_summary = content.get("pricing_summary", {})
    row_cursor += 2
    sheet[f"A{row_cursor}"] = _as_text(pricing_summary.get("label")) or "当前报价结构"
    sheet[f"B{row_cursor}"] = _as_text(pricing_summary.get("details"))
    sheet[f"C{row_cursor}"] = _as_text(pricing_summary.get("total"))

    row_cursor += 2
    sheet[f"A{row_cursor}"] = "阶段规划"
    row_cursor += 1
    sheet[f"A{row_cursor}"] = "阶段"
    sheet[f"B{row_cursor}"] = "板块"
    sheet[f"C{row_cursor}"] = "费用"
    sheet[f"D{row_cursor}"] = "周期"
    sheet[f"E{row_cursor}"] = "交付"
    for item in content.get("stage_plan", []):
        row_cursor += 1
        sheet[f"A{row_cursor}"] = _first_present(item, "stage", "name")
        sheet[f"B{row_cursor}"] = _first_present(item, "module")
        sheet[f"C{row_cursor}"] = _first_present(item, "price")
        sheet[f"D{row_cursor}"] = _first_present(item, "duration", "timeline")
        sheet[f"E{row_cursor}"] = _first_present(item, "deliverables")

    for section in content.get("service_sections", []):
        row_cursor += 2
        sheet[f"A{row_cursor}"] = _as_text(section.get("title"))

        rows = section.get("rows")
        if rows:
            row_cursor += 1
            headers = ["版块", "项目", "任务", "说明", "交付成果", "服务部门", "负责人"]
            for index, header in enumerate(headers, start=1):
                sheet.cell(row=row_cursor, column=index, value=header)

            for row in rows:
                row_cursor += 1
                sheet[f"A{row_cursor}"] = _first_present(row, "block")
                sheet[f"B{row_cursor}"] = _first_present(row, "project")
                sheet[f"C{row_cursor}"] = _first_present(row, "task")
                sheet[f"D{row_cursor}"] = _first_present(row, "description_bullets", "description")
                sheet[f"E{row_cursor}"] = _first_present(row, "deliverable", "deliverables")
                sheet[f"F{row_cursor}"] = _first_present(row, "department", "service_department")
                sheet[f"G{row_cursor}"] = _first_present(row, "owner")
            continue

        for item in section.get("items", []):
            row_cursor += 1
            sheet[f"A{row_cursor}"] = _as_text(item)

    closing_notes = content.get("closing_notes", [])
    if closing_notes:
        row_cursor += 2
        sheet[f"A{row_cursor}"] = "备注说明"
        for note in closing_notes:
            row_cursor += 1
            sheet[f"A{row_cursor}"] = _as_text(note)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _style_cell(
    value_cell,
    *,
    font: Font | None = None,
    fill: PatternFill | None = None,
    alignment: Alignment | None = None,
) -> None:
    if font is not None:
        value_cell.font = font
    if fill is not None:
        value_cell.fill = fill
    if alignment is not None:
        value_cell.alignment = alignment


def _apply_column_widths(sheet) -> None:
    widths = {
        "A": 18,
        "B": 22,
        "C": 20,
        "D": 36,
        "E": 28,
        "F": 18,
        "G": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _apply_table_header_style(sheet, header_row: int, column_count: int = 7) -> None:
    for column in range(1, column_count + 1):
        _style_cell(
            sheet.cell(row=header_row, column=column),
            font=HEADER_FONT,
            fill=HEADER_FILL,
            alignment=WRAP_TOP_ALIGNMENT,
        )


def _apply_table_body_style(sheet, start_row: int, end_row: int, column_count: int = 7) -> None:
    if end_row < start_row:
        return
    for row in range(start_row, end_row + 1):
        for column in range(1, column_count + 1):
            _style_cell(
                sheet.cell(row=row, column=column),
                alignment=WRAP_TOP_ALIGNMENT,
            )


def render_xlsx(
    content: dict[str, Any],
    output_path: Path,
    template_path: Path | None = None,
) -> Path:
    content = _normalized_content(content)
    if not _uses_stages(content):
        return _render_xlsx_legacy(content, output_path, template_path=template_path)

    if template_path and template_path.exists():
        workbook = load_workbook(template_path)
        sheet = workbook[workbook.sheetnames[0]]
    else:
        workbook = Workbook()
        sheet = workbook.active

    sheet.title = "报价方案"
    sheet["A1"] = _as_text(content.get("title"))
    sheet["A2"] = _as_text(content.get("subtitle"))
    _style_cell(sheet["A1"], font=TITLE_FONT, alignment=WRAP_TOP_ALIGNMENT)
    _style_cell(sheet["A2"], font=SUBTITLE_FONT, alignment=WRAP_TOP_ALIGNMENT)
    _apply_column_widths(sheet)

    row_cursor = 4
    sheet[f"A{row_cursor}"] = "项目概述"
    _style_cell(sheet[f"A{row_cursor}"], font=SECTION_FONT, alignment=WRAP_TOP_ALIGNMENT)
    for paragraph in content.get("overview_paragraphs", []):
        row_cursor += 1
        sheet[f"A{row_cursor}"] = _as_text(paragraph)
        _style_cell(sheet[f"A{row_cursor}"], alignment=WRAP_TOP_ALIGNMENT)

    pricing_summary = content.get("pricing_summary", {})
    row_cursor += 2
    sheet[f"A{row_cursor}"] = _as_text(pricing_summary.get("label")) or "当前报价结构"
    sheet[f"B{row_cursor}"] = _as_text(pricing_summary.get("details"))
    sheet[f"C{row_cursor}"] = _as_text(pricing_summary.get("total"))
    _style_cell(sheet[f"A{row_cursor}"], font=SECTION_FONT, alignment=WRAP_TOP_ALIGNMENT)
    _style_cell(sheet[f"B{row_cursor}"], alignment=WRAP_TOP_ALIGNMENT)
    _style_cell(sheet[f"C{row_cursor}"], alignment=WRAP_TOP_ALIGNMENT)

    headers = [_top_level_label(content), "项目", "任务", "说明", "交付物", "服务部门", "负责人"]
    header_frozen = False
    for section in build_stage_table_sections(content):
        row_cursor += 2
        sheet[f"A{row_cursor}"] = section["stage_name"]
        sheet[f"B{row_cursor}"] = "阶段费用"
        sheet[f"C{row_cursor}"] = _as_text(section.get("stage_price")) or "待补充"
        sheet[f"D{row_cursor}"] = "阶段周期"
        sheet[f"E{row_cursor}"] = _as_text(section.get("stage_duration")) or "待确认"
        for cell_ref in (f"A{row_cursor}", f"B{row_cursor}", f"C{row_cursor}", f"D{row_cursor}", f"E{row_cursor}"):
            _style_cell(sheet[cell_ref], alignment=WRAP_TOP_ALIGNMENT)
        _style_cell(sheet[f"A{row_cursor}"], font=SECTION_FONT, alignment=WRAP_TOP_ALIGNMENT)
        _style_cell(sheet[f"B{row_cursor}"], font=SECTION_FONT, alignment=WRAP_TOP_ALIGNMENT)
        _style_cell(sheet[f"D{row_cursor}"], font=SECTION_FONT, alignment=WRAP_TOP_ALIGNMENT)

        row_cursor += 1
        sheet[f"A{row_cursor}"] = "阶段交付物"
        sheet[f"B{row_cursor}"] = _as_text(section.get("stage_deliverables_text"))
        _style_cell(sheet[f"A{row_cursor}"], font=SECTION_FONT, alignment=WRAP_TOP_ALIGNMENT)
        _style_cell(sheet[f"B{row_cursor}"], alignment=WRAP_TOP_ALIGNMENT)

        row_cursor += 1
        header_row = row_cursor
        for index, header in enumerate(headers, start=1):
            sheet.cell(row=header_row, column=index, value=header)
        _apply_table_header_style(sheet, header_row)
        if not header_frozen:
            sheet.freeze_panes = f"A{header_row + 1}"
            header_frozen = True

        data_start_row = header_row + 1
        current_row = data_start_row
        for row in section["rows"]:
            sheet[f"A{current_row}"] = row["stage_name"] if row["show_stage"] else ""
            sheet[f"B{current_row}"] = row["project_name"] if row["show_project"] else ""
            sheet[f"C{current_row}"] = row["task_name"]
            sheet[f"D{current_row}"] = row["description_text"]
            sheet[f"E{current_row}"] = row["deliverables_text"] if row["show_deliverables"] else ""
            sheet[f"F{current_row}"] = row["department_text"]
            sheet[f"G{current_row}"] = row["owner_text"]
            current_row += 1

        data_end_row = current_row - 1
        _apply_table_body_style(sheet, data_start_row, data_end_row)
        if section["rows"]:
            if section["rows"][0]["stage_rowspan"] > 1:
                sheet.merge_cells(
                    start_row=data_start_row,
                    start_column=1,
                    end_row=data_end_row,
                    end_column=1,
                )
            cursor = data_start_row
            for row in section["rows"]:
                if row["show_project"] and row["project_rowspan"] > 1:
                    end_row = cursor + row["project_rowspan"] - 1
                    sheet.merge_cells(start_row=cursor, start_column=2, end_row=end_row, end_column=2)
                    sheet.merge_cells(start_row=cursor, start_column=5, end_row=end_row, end_column=5)
                cursor += 1

        row_cursor = data_end_row

    closing_notes = content.get("closing_notes", [])
    if closing_notes:
        row_cursor += 2
        sheet[f"A{row_cursor}"] = "备注说明"
        _style_cell(sheet[f"A{row_cursor}"], font=SECTION_FONT, alignment=WRAP_TOP_ALIGNMENT)
        for note in closing_notes:
            row_cursor += 1
            sheet[f"A{row_cursor}"] = _as_text(note)
            _style_cell(sheet[f"A{row_cursor}"], alignment=WRAP_TOP_ALIGNMENT)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _load_reportlab() -> dict[str, Any]:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "缺少 Python 依赖 reportlab。请先运行: python3 -m pip install --target /Users/a123/.openclaw/skills/quote-skill/python reportlab"
        ) from exc

    return {
        "colors": colors,
        "A4": A4,
        "getSampleStyleSheet": getSampleStyleSheet,
        "pdfmetrics": pdfmetrics,
        "UnicodeCIDFont": UnicodeCIDFont,
        "Paragraph": Paragraph,
        "SimpleDocTemplate": SimpleDocTemplate,
        "Spacer": Spacer,
        "Table": Table,
        "TableStyle": TableStyle,
    }


def _render_pdf_reportlab(content: dict[str, Any], output_path: Path) -> Path:
    deps = _load_reportlab()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    deps["pdfmetrics"].registerFont(deps["UnicodeCIDFont"]("STSong-Light"))
    document = deps["SimpleDocTemplate"](str(output_path), pagesize=deps["A4"])
    styles = deps["getSampleStyleSheet"]()
    styles["Title"].fontName = "STSong-Light"
    styles["Heading2"].fontName = "STSong-Light"
    styles["Heading3"].fontName = "STSong-Light"
    styles["BodyText"].fontName = "STSong-Light"

    story = [deps["Paragraph"](_as_text(content.get("title")), styles["Title"])]
    subtitle = _as_text(content.get("subtitle"))
    if subtitle:
        story.append(deps["Paragraph"](subtitle, styles["Heading3"]))
    story.append(deps["Spacer"](1, 12))

    for paragraph in content.get("overview_paragraphs", []):
        story.append(deps["Paragraph"](_as_text(paragraph), styles["BodyText"]))
        story.append(deps["Spacer"](1, 8))

    pricing_summary_lines = _pricing_summary_lines(content.get("pricing_summary", {}))
    if pricing_summary_lines:
        story.append(deps["Paragraph"]("当前报价结构", styles["Heading2"]))
        for line in pricing_summary_lines:
            story.append(deps["Paragraph"](line, styles["BodyText"]))
        story.append(deps["Spacer"](1, 12))

    if _uses_stages(content):
        stage_rows = [["阶段", "项目", "任务", "说明", "交付物", "服务部门", "负责人"]]
        for section in build_stage_table_sections(content):
            for row in section["rows"]:
                stage_rows.append(
                    [
                        row["stage_name"] if row["show_stage"] else "",
                        row["project_name"] if row["show_project"] else "",
                        row["task_name"],
                        row["description_text"],
                        row["deliverables_text"] if row["show_deliverables"] else "",
                        row["department_text"],
                        row["owner_text"],
                    ]
                )
        if len(stage_rows) > 1:
            story.append(deps["Paragraph"]("项目阶段规划", styles["Heading2"]))
            stage_table = deps["Table"](stage_rows, repeatRows=1)
            stage_table.setStyle(
                deps["TableStyle"](
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), deps["colors"].HexColor("#f2e7da")),
                        ("GRID", (0, 0), (-1, -1), 0.5, deps["colors"].HexColor("#ddd2c5")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                    ]
                )
            )
            story.append(stage_table)
            story.append(deps["Spacer"](1, 12))

    for section in content.get("service_sections", []):
        title = _as_text(section.get("title"))
        if title:
            story.append(deps["Paragraph"](title, styles["Heading2"]))

        rows = section.get("rows")
        if rows:
            service_rows = [["版块", "项目", "任务", "说明", "交付成果", "服务部门", "负责人"]]
            for row in rows:
                service_rows.append(
                    [
                        _first_present(row, "block"),
                        _first_present(row, "project"),
                        _first_present(row, "task"),
                        _first_present(row, "description_bullets", "description"),
                        _first_present(row, "deliverable", "deliverables"),
                        _first_present(row, "department", "service_department"),
                        _first_present(row, "owner"),
                    ]
                )
            service_table = deps["Table"](service_rows, repeatRows=1)
            service_table.setStyle(
                deps["TableStyle"](
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), deps["colors"].HexColor("#f2e7da")),
                        ("GRID", (0, 0), (-1, -1), 0.5, deps["colors"].HexColor("#ddd2c5")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                    ]
                )
            )
            story.append(service_table)
            story.append(deps["Spacer"](1, 12))
            continue

        for item in section.get("items", []):
            story.append(deps["Paragraph"](f"- {_as_text(item)}", styles["BodyText"]))
        story.append(deps["Spacer"](1, 12))

    closing_notes = content.get("closing_notes", [])
    if closing_notes:
        story.append(deps["Paragraph"]("备注说明", styles["Heading2"]))
        for note in closing_notes:
            story.append(deps["Paragraph"](_as_text(note), styles["BodyText"]))

    document.build(story)
    return output_path


def _resolve_node_modules_path(node_bin: Path) -> Path:
    if node_bin.parent.name == "bin":
        return node_bin.parent.parent / "node_modules"
    return node_bin.parent / "node_modules"


def _format_html_to_pdf_failure(
    prefix: str,
    *,
    stdout: str | None = None,
    stderr: str | None = None,
) -> str:
    details = [prefix]
    if stdout:
        details.append(f"stdout:\n{stdout.strip()}")
    if stderr:
        details.append(f"stderr:\n{stderr.strip()}")
    return "\n\n".join(details)


def render_pdf_from_html(
    html_path: Path,
    output_path: Path,
    node_bin: Path | None = None,
    html_to_pdf_script: Path | None = None,
) -> Path:
    node_bin = node_bin or DEFAULT_NODE_BIN
    html_to_pdf_script = html_to_pdf_script or DEFAULT_HTML_TO_PDF_SCRIPT
    html_path = Path(html_path)
    output_path = Path(output_path)

    if not html_path.exists():
        raise FileNotFoundError(f"找不到要转 PDF 的 HTML 文件：{html_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["NODE_PATH"] = str(_resolve_node_modules_path(node_bin))

    try:
        result = subprocess.run(
            [
                str(node_bin),
                str(html_to_pdf_script),
                str(html_path),
                str(output_path),
            ],
            check=True,
            env=env,
            capture_output=True,
            text=True,
        )
    except subprocess.CalledProcessError as exc:
        raise RuntimeError(
            _format_html_to_pdf_failure(
                f"浏览器转 PDF 失败：{html_path} -> {output_path}",
                stdout=exc.stdout,
                stderr=exc.stderr,
            )
        ) from exc

    if not output_path.exists():
        raise RuntimeError(
            _format_html_to_pdf_failure(
                f"浏览器转 PDF 已执行，但没有生成 PDF 文件：{output_path}",
                stdout=getattr(result, "stdout", None),
                stderr=getattr(result, "stderr", None),
            )
        )
    return output_path


def _render_pdf_via_browser(
    content: dict[str, Any],
    output_path: Path,
    template_path: Path,
    node_bin: Path,
    html_to_pdf_script: Path,
) -> Path:
    with tempfile.TemporaryDirectory(prefix="quote-pdf-html-") as tmp_dir:
        staged_html_path = Path(tmp_dir) / f"{output_path.stem}.html"
        render_html(content, template_path=template_path, output_path=staged_html_path)
        return render_pdf_from_html(
            staged_html_path,
            output_path,
            node_bin=node_bin,
            html_to_pdf_script=html_to_pdf_script,
        )


def render_pdf(
    content: dict[str, Any],
    output_path: Path,
    template_path: Path | None = None,
    node_bin: Path | None = None,
    html_to_pdf_script: Path | None = None,
) -> Path:
    content = _normalized_content(content)
    try:
        if _uses_stages(content):
            node_bin = node_bin or DEFAULT_NODE_BIN
            html_to_pdf_script = html_to_pdf_script or DEFAULT_HTML_TO_PDF_SCRIPT
            if node_bin and Path(node_bin).exists() and html_to_pdf_script.exists():
                return _render_pdf_via_browser(
                    content,
                    output_path,
                    template_path=template_path or DEFAULT_HTML_TEMPLATE_PATH,
                    node_bin=Path(node_bin),
                    html_to_pdf_script=html_to_pdf_script,
                )
        return _render_pdf_reportlab(content, output_path)
    except RuntimeError as exc:
        if "reportlab" in str(exc).lower():
            raise
        return _render_pdf_reportlab(content, output_path)
