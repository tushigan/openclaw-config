from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
import subprocess
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from quote_skill.renderers import render_html, render_pdf, render_pdf_from_html, render_xlsx

NODE_BIN = Path(
    "/Users/a123/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node"
)
HTML_TO_PDF_SCRIPT = (
    Path(__file__).resolve().parents[1] / "scripts" / "render_html_to_pdf.mjs"
)
HTML_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1] / "assets" / "templates" / "quote.html"
)


def sample_quote_content(case_type: str = "散案") -> dict:
    content = {
        "title": "品牌升级报价方案",
        "subtitle": "阶段化服务测试",
        "case_type": case_type,
        "overview_paragraphs": ["测试概述。"],
        "pricing_summary": {
            "label": "当前报价结构",
            "details": "按阶段执行",
            "total": "待补充",
        },
        "stages": [
            {
                "stage_name": "第一阶段",
                "stage_price": "待补充",
                "stage_duration": "2 周",
                "stage_deliverables": ["阶段交付 A"],
                "projects": [
                    {
                        "project_name": "项目一",
                        "project_deliverables": ["项目一交付"],
                        "service_catalog_refs": ["catalog-1", "catalog-2"],
                        "tasks": [
                            {
                                "task_name": "任务一",
                                "description_bullets": ["说明一"],
                                "service_catalog_refs": ["catalog-1"],
                            },
                            {
                                "task_name": "任务二",
                                "description_bullets": ["说明二"],
                                "service_catalog_refs": ["catalog-2"],
                            },
                        ],
                    },
                    {
                        "project_name": "项目二",
                        "project_deliverables": ["项目二交付"],
                        "service_catalog_refs": ["catalog-3", "catalog-4"],
                        "tasks": [
                            {
                                "task_name": "任务三",
                                "description_bullets": ["说明三"],
                                "service_catalog_refs": ["catalog-3"],
                            },
                            {
                                "task_name": "任务四",
                                "description_bullets": ["说明四"],
                                "service_catalog_refs": ["catalog-4"],
                            },
                        ],
                    },
                ],
            }
        ],
        "closing_notes": ["报价有效期 15 天。", "排期以合同签署时间为准。"],
    }
    if case_type == "全案":
        content["annual_context"] = {
            "business_goal": "建立全年品牌经营节奏",
            "service_period": "2026 年全年",
            "service_boundary": "品牌经营与包装升级支持",
            "core_service_modules": ["品牌策略", "包装体系"],
            "priority": "高",
        }
    return content


class RendererTests(unittest.TestCase):
    def test_render_html_includes_rowspan_for_stage_and_project(self):
        quote_content = sample_quote_content()

        html = render_html(quote_content, template_path=HTML_TEMPLATE_PATH)

        self.assertIn("品牌升级报价方案", html)
        self.assertIn("第一阶段", html)
        self.assertIn('rowspan="4"', html)
        self.assertIn('rowspan="2"', html)

    def test_render_html_uses_service_module_label_for_full_case(self):
        quote_content = sample_quote_content(case_type="全案")

        html = render_html(quote_content, template_path=HTML_TEMPLATE_PATH)

        self.assertIn("<th>服务模块</th><th>项目</th><th>任务</th>", html)
        self.assertNotIn("<th>阶段</th><th>项目</th><th>任务</th>", html)

    def test_render_xlsx_writes_title_and_default_styles(self):
        quote_content = sample_quote_content()

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "quote.xlsx"
            template_path = (
                Path(__file__).resolve().parents[1] / "assets" / "templates" / "quote.xlsx"
            )

            render_xlsx(quote_content, output_path, template_path=template_path)

            workbook = load_workbook(output_path)
            values = [
                cell
                for row in workbook["报价方案"].iter_rows(values_only=True)
                for cell in row
                if cell not in (None, "")
            ]
            sheet = workbook["报价方案"]
            merged_ranges = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
            header_values = ["阶段", "项目", "任务", "说明", "交付物", "服务部门", "负责人"]
            header_row = next(
                row_idx
                for row_idx in range(1, sheet.max_row + 1)
                if [sheet.cell(row=row_idx, column=col).value for col in range(1, 8)] == header_values
            )
            self.assertIn("报价方案", workbook.sheetnames)
            self.assertEqual(sheet["A1"].value, "品牌升级报价方案")
            self.assertIn("第一阶段", values)
            self.assertIn("项目一", values)
            self.assertIn("任务一", values)
            self.assertIn("报价有效期 15 天。", values)
            self.assertTrue(any(cell_range.startswith("A") for cell_range in merged_ranges))
            self.assertTrue(any(cell_range.startswith("B") for cell_range in merged_ranges))
            self.assertEqual(sheet.freeze_panes, f"A{header_row + 1}")
            self.assertEqual(sheet.cell(row=header_row, column=1).fill.fgColor.rgb, "00F2E7DA")
            self.assertTrue(sheet.cell(row=header_row, column=1).font.bold)
            self.assertTrue(sheet.cell(row=header_row, column=1).alignment.wrap_text)
            self.assertEqual(sheet.cell(row=header_row, column=1).alignment.vertical, "top")
            self.assertGreater(sheet.column_dimensions["D"].width, 20)
            self.assertTrue(sheet.cell(row=header_row + 1, column=4).alignment.wrap_text)
            self.assertEqual(sheet.cell(row=header_row + 1, column=4).alignment.vertical, "top")

    def test_render_pdf_calls_browser_html_to_pdf_pipeline(self):
        quote_content = sample_quote_content()
        node_bin = Path("/tmp/fake-node")
        html_to_pdf_script = Path("/tmp/render_html_to_pdf.mjs")

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "quote.pdf"
            html_path = Path(tmp) / "quote.html"
            render_html(quote_content, template_path=HTML_TEMPLATE_PATH, output_path=html_path)
            recorded: dict[str, object] = {}

            def fake_run(command, check, env, capture_output, text):
                recorded["command"] = command
                recorded["check"] = check
                recorded["env"] = env
                recorded["capture_output"] = capture_output
                recorded["text"] = text
                Path(command[3]).write_bytes(b"%PDF-1.4\n%fake\n")

            with patch("quote_skill.renderers.subprocess.run", side_effect=fake_run):
                render_pdf_from_html(
                    html_path,
                    output_path,
                    node_bin=node_bin,
                    html_to_pdf_script=html_to_pdf_script,
                )

            command = recorded["command"]
            self.assertEqual(command[0], str(node_bin))
            self.assertEqual(command[1], str(html_to_pdf_script))
            self.assertEqual(command[2], str(html_path))
            self.assertEqual(command[3], str(output_path))
            self.assertTrue(output_path.is_file())
            self.assertEqual(recorded["env"]["NODE_PATH"], str(node_bin.parent / "node_modules"))
            html_text = html_path.read_text(encoding="utf-8")
            self.assertIn("品牌升级报价方案", html_text)
            self.assertIn("第一阶段", html_text)
            self.assertIn("任务一", html_text)

    def test_render_pdf_uses_temp_html_instead_of_output_dir_by_default(self):
        quote_content = sample_quote_content()

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "quote.pdf"

            with patch("quote_skill.renderers.render_pdf_from_html") as pdf_from_html:
                render_pdf(
                    quote_content,
                    output_path,
                    template_path=HTML_TEMPLATE_PATH,
                    node_bin=NODE_BIN,
                    html_to_pdf_script=HTML_TO_PDF_SCRIPT,
                )

            html_input_path = Path(pdf_from_html.call_args.args[0])
            self.assertNotEqual(html_input_path.parent, output_path.parent)
            self.assertFalse((output_path.parent / "quote.staged.html").exists())

    def test_render_pdf_from_html_surfaces_stdout_and_stderr_on_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            html_path = Path(tmp) / "quote.html"
            output_path = Path(tmp) / "quote.pdf"
            html_path.write_text("<html><body>test</body></html>", encoding="utf-8")

            error = subprocess.CalledProcessError(
                1,
                ["node", "render_html_to_pdf.mjs"],
                output="stdout text",
                stderr="stderr text",
            )

            with patch("quote_skill.renderers.subprocess.run", side_effect=error):
                with self.assertRaisesRegex(RuntimeError, "stderr text"):
                    render_pdf_from_html(
                        html_path,
                        output_path,
                        node_bin=NODE_BIN,
                        html_to_pdf_script=HTML_TO_PDF_SCRIPT,
                    )

    def test_render_pdf_real_browser_pipeline_generates_pdf(self):
        quote_content = sample_quote_content()

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "quote.pdf"

            render_pdf(
                quote_content,
                output_path,
                template_path=HTML_TEMPLATE_PATH,
                node_bin=NODE_BIN,
                html_to_pdf_script=HTML_TO_PDF_SCRIPT,
            )

            self.assertTrue(output_path.is_file())
            self.assertGreater(output_path.stat().st_size, 0)
            self.assertGreaterEqual(len(PdfReader(str(output_path)).pages), 1)
            self.assertFalse((output_path.parent / "quote.staged.html").exists())


if __name__ == "__main__":
    unittest.main()
